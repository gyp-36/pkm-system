"""M3 file notes, content exports, and account-scoped web import drafts."""

import hashlib
import http.client
import ipaddress
import io
import json
import os
import re
import socket
import ssl
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.robotparser import RobotFileParser
from urllib.parse import quote, urljoin, urlparse, urlsplit, urlunparse, urlunsplit

import httpx
import jwt
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from fastapi import APIRouter, File, Form, HTTPException, Query, Request, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field, HttpUrl
from pypdf import PdfReader
from PIL import Image, ImageOps
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from starlette.responses import Response

from app.auth.auth import Db, UserId
from app.assistant.content_analysis import ContentAnalysis, analyze_content, generate_link_note_content
from app.core import object_storage
from app.core.enums import AuditAction, AuditActor, AuditEntityType
from app.core.lifecycle import record_event, record_revision
from app.core.models import FileIngestJob, LinkDraft, Note, NoteFileVersion, NoteTextBlock, Notebook, NoteTag
from app.core.rate_limit import check_limit
from app.knowledge.notes import download_response, owned_note, queue_index, safe_filename, validate_categories
from app.knowledge.file_types import ALLOWED_EXTENSIONS, IMAGE_EXTENSIONS, MIME_BY_EXT, PIL_FORMAT_BY_EXT
from app.knowledge.text_safety import sanitize_extracted_text

router = APIRouter(tags=["M3 notes"])
MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_WEB_BYTES = 3 * 1024 * 1024
FILE_ROOT = Path(os.getenv("NOTE_FILE_ROOT", "/data/note-files"))


def _file_path(storage_key: str) -> Path:
    if not re.fullmatch(r"[0-9a-f-]{36}", storage_key):
        raise HTTPException(status_code=500, detail="文件存储记录无效")
    return FILE_ROOT / storage_key[:2] / storage_key


def _file_bytes(row: NoteFileVersion) -> bytes:
    if row.storage_backend == "s3":
        try:
            return object_storage.get(row.storage_key)
        except Exception:
            raise HTTPException(status_code=410, detail="文件内容不可用") from None
    path = _file_path(row.storage_key)
    if not path.is_file():
        raise HTTPException(status_code=410, detail="文件内容不可用")
    try:
        return path.read_bytes()
    except OSError:
        raise HTTPException(status_code=410, detail="文件内容不可用") from None


def _onlyoffice_download_url(source: str | None) -> str:
    """将 ONLYOFFICE 面向浏览器的下载地址映射到容器内地址。"""
    public = urlparse(os.getenv("ONLYOFFICE_PUBLIC_URL", ""))
    internal = urlparse(os.getenv("ONLYOFFICE_INTERNAL_URL", "http://documentserver"))
    parsed = urlparse(source or "")
    if not parsed.scheme or not parsed.hostname or parsed.username or parsed.password or parsed.fragment:
        raise HTTPException(status_code=400, detail="文档服务回调地址无效")

    def same_origin(left, right) -> bool:
        return (
            left.scheme == right.scheme
            and left.hostname == right.hostname
            and left.port == right.port
        )

    if same_origin(parsed, internal):
        return source or ""
    if public.hostname and same_origin(parsed, public):
        return urlunparse(parsed._replace(scheme=internal.scheme, netloc=internal.netloc))
    raise HTTPException(status_code=400, detail="文档服务回调地址无效")


def _queue_file_ingest(db: Db, note: Note, row: NoteFileVersion) -> None:
    existing = db.scalar(select(FileIngestJob).where(FileIngestJob.file_version_id == row.id))
    if existing is None:
        db.add(FileIngestJob(user_id=note.user_id, note_id=note.id, file_version_id=row.id, status="pending"))


def _latest_version(db: Db, note: Note) -> NoteFileVersion | None:
    return db.scalar(select(NoteFileVersion).where(NoteFileVersion.note_id == note.id, NoteFileVersion.user_id == note.user_id).order_by(NoteFileVersion.version.desc()).limit(1))


def find_duplicate_file(db: Db, user_id: UserId, sha256: str) -> tuple[Note, NoteFileVersion] | None:
    match = db.execute(
        select(Note, NoteFileVersion)
        .join(NoteFileVersion, (NoteFileVersion.note_id == Note.id) & (NoteFileVersion.user_id == Note.user_id))
        .where(Note.user_id == user_id, Note.deleted_at.is_(None), NoteFileVersion.sha256 == sha256)
        .order_by(Note.updated_at.desc(), Note.id)
        .limit(1)
    ).first()
    return (match[0], match[1]) if match else None


class FileHashCheck(BaseModel):
    sha256: str = Field(pattern=r"^[0-9a-fA-F]{64}$")


def duplicate_file_error(note: Note, row: NoteFileVersion) -> HTTPException:
    return HTTPException(status_code=409, detail={
        "code": "duplicate_file",
        "note_id": str(note.id),
        "title": note.title,
        "filename": row.filename,
    })


def file_meta(row: NoteFileVersion | None) -> dict | None:
    if row is None:
        return None
    return {"filename": row.filename, "extension": row.extension, "media_type": row.media_type, "size_bytes": row.size_bytes, "file_version": row.version, "sha256": row.sha256, "extraction_status": row.extraction_status, "extraction_error": row.extraction_error, "extraction_fingerprint": row.extraction_fingerprint}


def file_content_disposition(disposition: str, filename: str) -> str:
    safe_name = safe_filename(filename, "file")
    ascii_name = safe_name.encode("ascii", "ignore").decode("ascii") or "file"
    return f'{disposition}; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(safe_name, safe="")}'


def extract_text(extension: str, content: bytes) -> tuple[str, list[tuple[int, int, dict]]]:
    pieces: list[str] = []
    blocks: list[tuple[int, int, dict]] = []
    cursor = 0

    def append(text: str, locator: dict) -> None:
        nonlocal cursor
        text = sanitize_extracted_text(text.strip())
        if not text:
            return
        if pieces:
            pieces.append("\n\n")
            cursor += 2
        start = cursor
        pieces.append(text)
        cursor += len(text)
        blocks.append((start, cursor, locator))

    try:
        if extension == "md":
            text = sanitize_extracted_text(content.decode("utf-8-sig")[:1_000_000])
            return text, ([(0, len(text), {"kind": "markdown"})] if text else [])
        if extension == "docx":
            doc = Document(io.BytesIO(content))
            paragraph_number = 0
            table_number = 0
            heading_path: list[str] = []
            for element in doc.element.body.iterchildren():
                if element.tag.endswith("}p"):
                    paragraph_number += 1
                    paragraph = next((item for item in doc.paragraphs if item._p is element), None)
                    text = paragraph.text if paragraph else ""
                    style_name = paragraph.style.name if paragraph and paragraph.style else ""
                    heading = re.match(r"^(?:Heading|标题)\s*(\d+)$", style_name, re.IGNORECASE)
                    if heading and text.strip():
                        level = max(1, min(6, int(heading.group(1))))
                        heading_path = heading_path[: level - 1]
                        heading_path.append(text.strip())
                        kind = "heading"
                    else:
                        kind = "paragraph"
                    append(text, {"kind": kind, "paragraph": paragraph_number, "heading_path": list(heading_path)})
                elif element.tag.endswith("}tbl"):
                    table_number += 1
                    table = next((item for item in doc.tables if item._tbl is element), None)
                    if table:
                        for row_number, row in enumerate(table.rows, start=1):
                            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                            rendered = " | ".join(value for value in values if value)
                            append(rendered, {"kind": "table_row", "table": table_number, "row": row_number, "heading_path": list(heading_path)})
            return "".join(pieces)[:1_000_000], blocks
        if extension == "xlsx":
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            for sheet in book.worksheets:
                append(f"工作表：{sheet.title}", {"kind": "sheet", "sheet": sheet.title})
                for row_number, row in enumerate(sheet.iter_rows(), start=1):
                    values = [f"{cell.coordinate}: {cell.value}" for cell in row if cell.value is not None]
                    if values:
                        append(" | ".join(values), {"kind": "sheet_row", "sheet": sheet.title, "row": row_number, "cells": [cell.coordinate for cell in row if cell.value is not None]})
            return "".join(pieces)[:1_000_000], blocks
        if extension == "pdf":
            pdf = PdfReader(io.BytesIO(content), strict=False)
            for ordinal, page in enumerate(pdf.pages, start=1):
                append(page.extract_text() or "", {"kind": "page", "page": ordinal})
            return "".join(pieces)[:1_000_000], blocks
    except Exception:
        return "", []
    return "", []


def persist_file(db: Db, note: Note, filename: str, content: bytes, extension: str) -> NoteFileVersion:
    key = str(uuid.uuid4())
    object_storage.put(key, content, MIME_BY_EXT[extension])
    current = _latest_version(db, note)
    row = NoteFileVersion(
        user_id=note.user_id, note_id=note.id, version=(current.version + 1 if current else 1),
        filename=safe_filename(Path(filename).name, "file")[:255], extension=extension, media_type=MIME_BY_EXT[extension],
        storage_key=key, size_bytes=len(content), sha256=hashlib.sha256(content).hexdigest(), storage_backend="s3", extraction_status="pending",
    )
    db.add(row)
    note.content_kind = "markdown" if extension == "md" else extension
    note.body_md = ""
    note.content_version += 1
    queue_index(db, note)
    db.flush()
    _queue_file_ingest(db, note, row)
    return row


def _check_file(ext: str, content: bytes) -> None:
    if not content:
        raise HTTPException(status_code=422, detail="文件内容为空")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="文件最大支持 25 MB")
    signatures = {"pdf": content.startswith(b"%PDF-"), "docx": content.startswith(b"PK\x03\x04"), "xlsx": content.startswith(b"PK\x03\x04")}
    if ext in signatures and not signatures[ext]:
        raise HTTPException(status_code=415, detail="文件内容与扩展名不匹配或文件已损坏")
    if ext in {"doc", "xls"}:
        # 旧版 Office 二进制格式会作为原始文件保留，直到用户明确执行转换。
        if not content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raise HTTPException(status_code=415, detail="旧版 Office 文件内容无效")
    try:
        if ext == "md":
            content.decode("utf-8-sig")
        elif ext == "docx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if archive.testzip() is not None:
                    raise ValueError("corrupt zip")
            Document(io.BytesIO(content))
        elif ext == "xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if archive.testzip() is not None:
                    raise ValueError("corrupt zip")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            if not workbook.sheetnames:
                raise ValueError("no worksheets")
        elif ext == "pdf":
            pdf = PdfReader(io.BytesIO(content), strict=True)
            if not pdf.pages:
                raise ValueError("no pages")
        elif ext in IMAGE_EXTENSIONS:
            image = Image.open(io.BytesIO(content))
            if image.format not in PIL_FORMAT_BY_EXT[ext]:
                raise ValueError("image format does not match extension")
            image.verify()
    except Exception:
        raise HTTPException(status_code=415, detail="文件格式无效或内容已损坏") from None


def convert_legacy_office(filename: str, extension: str, storage_key: str, media_type: str) -> tuple[str, bytes]:
    office = os.getenv("ONLYOFFICE_INTERNAL_URL", "").rstrip("/")
    secret = os.getenv("ONLYOFFICE_JWT_SECRET", "")
    if not office or not secret:
        raise HTTPException(status_code=503, detail="旧版 Office 转换需要先启动 ONLYOFFICE 文档服务")
    key = str(uuid.uuid4())
    expiry = int((datetime.now(timezone.utc) + timedelta(minutes=10)).timestamp())
    link_token = jwt.encode({"key": storage_key, "media_type": media_type, "exp": expiry}, os.getenv("PKM_FILE_LINK_SECRET", secret), algorithm="HS256")
    app_url = os.getenv("PUBLIC_APP_URL", "http://api:8000").rstrip("/")
    file_url = f"{app_url}/v1/files/{storage_key}?token={link_token}"
    target = "docx" if extension == "doc" else "xlsx"
    payload = {"async": False, "filetype": extension, "key": key, "outputtype": target, "title": filename, "url": file_url}
    payload["token"] = jwt.encode(payload.copy(), secret, algorithm="HS256")
    try:
        with httpx.Client(timeout=httpx.Timeout(90, connect=5)) as client:
            response = client.post(f"{office}/converter", json=payload, headers={"Accept": "application/json"})
            response.raise_for_status()
            result = response.json()
            result_url = result.get("fileUrl")
            allowed = urlparse(office)
            actual = urlparse(result_url or "")
            if result.get("error", 0) != 0 or not result.get("endConvert") or actual.hostname != allowed.hostname or actual.port != allowed.port:
                raise ValueError("conversion failed")
            with client.stream("GET", result_url) as converted:
                converted.raise_for_status()
                output = bytearray()
                for chunk in converted.iter_bytes():
                    output.extend(chunk)
                    if len(output) > MAX_UPLOAD_BYTES:
                        raise ValueError("converted file too large")
    except (httpx.HTTPError, ValueError, KeyError):
        raise HTTPException(status_code=422, detail="ONLYOFFICE 无法转换此文件") from None
    if len(output) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="转换后的文件超过 25 MB")
    return target, output


def note_payload(db: Db, note: Note) -> dict:
    from app.knowledge.notes import note_json
    data = note_json(db, note)
    data["content_kind"] = note.content_kind or "markdown"
    data["file"] = file_meta(_latest_version(db, note))
    current = _latest_version(db, note)
    data["extraction_status"] = current.extraction_status if current else ("available" if note.body_md.strip() else "empty")
    return data


@router.post("/v1/notes/upload", status_code=201)
async def upload_note(db: Db, user_id: UserId, file: UploadFile = File(...), sha256: str = Form(...), notebook_id: uuid.UUID | None = None, convert_legacy: bool = False) -> dict:
    filename = Path(file.filename or "upload").name
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="支持 md、docx、xlsx、pdf、png、jpg、jpeg、gif、bmp、webp、tif、tiff、ico、doc、xls")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    _check_file(ext, content)
    file_sha256 = hashlib.sha256(content).hexdigest()
    if not re.fullmatch(r"[0-9a-fA-F]{64}", sha256) or sha256.lower() != file_sha256:
        raise HTTPException(status_code=422, detail="文件哈希校验失败，请重新选择文件")
    duplicate = find_duplicate_file(db, user_id, file_sha256)
    if duplicate:
        raise duplicate_file_error(*duplicate)
    if ext in {"doc", "xls"} and not convert_legacy:
        raise HTTPException(status_code=409, detail="上传旧版 Office 文件前需确认转换为 DOCX/XLSX，原件会保留供下载")
    validate_categories(db, user_id, notebook_id, [])
    note = Note(user_id=user_id, notebook_id=notebook_id, title=Path(filename).stem[:240] or "未命名文件", body_md="", version=1, content_version=1, content_kind=ext)
    db.add(note)
    db.flush()
    if ext in {"doc", "xls"}:
        original = persist_file(db, note, filename, content, ext)
        db.flush()
        try:
            converted_ext, converted = convert_legacy_office(filename, ext, original.storage_key, original.media_type)
        except HTTPException:
            object_storage.delete(original.storage_key)
            raise
        original.extraction_status = "preserved"
        persist_file(db, note, f"{Path(filename).stem}.{converted_ext}", converted, converted_ext)
    else:
        persist_file(db, note, filename, content, ext)
    record_revision(db, note, [])
    record_event(db, user_id, AuditAction.CREATE, AuditEntityType.NOTE, note.id, entity_version=note.version, details={"fields": ["content_kind"]})
    db.commit()
    db.refresh(note)
    return note_payload(db, note)


@router.post("/v1/notes/file-duplicate-check")
def check_file_duplicate(body: FileHashCheck, db: Db, user_id: UserId) -> dict:
    duplicate = find_duplicate_file(db, user_id, body.sha256.lower())
    if duplicate is None:
        return {"exists": False}
    note, row = duplicate
    return {"exists": True, "note_id": str(note.id), "title": note.title, "filename": row.filename}


@router.get("/v1/notes/{note_id}/file")
def get_note_file(note_id: uuid.UUID, db: Db, user_id: UserId, version: int | None = None, download: bool = False) -> Response:
    note = owned_note(db, note_id, user_id)
    query = select(NoteFileVersion).where(NoteFileVersion.note_id == note.id, NoteFileVersion.user_id == user_id)
    if version is not None:
        query = query.where(NoteFileVersion.version == version)
    row = db.scalar(query.order_by(NoteFileVersion.version.desc()).limit(1))
    if row is None:
        raise HTTPException(status_code=404, detail="文件版本不存在")
    disposition = "attachment" if download else "inline"
    content = _file_bytes(row)
    media_type = row.media_type
    if not download and row.extension in {"tif", "tiff", "ico"}:
        # Browsers do not consistently render TIFF or ICO in an <img>; keep
        # downloads untouched and normalize only the inline preview response.
        with Image.open(io.BytesIO(content)) as source:
            image = ImageOps.exif_transpose(source)
            image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
            preview = io.BytesIO()
            image.save(preview, format="PNG")
        content = preview.getvalue()
        media_type = "image/png"
    return Response(content, media_type=media_type, headers={"Content-Disposition": file_content_disposition(disposition, row.filename)})


@router.put("/v1/notes/{note_id}/file")
async def replace_note_file(note_id: uuid.UUID, db: Db, user_id: UserId, file: UploadFile = File(...), version: int = Query(ge=1)) -> dict:
    note = owned_note(db, note_id, user_id, lock=True)
    if note.version != version:
        raise HTTPException(status_code=409, detail="笔记已有新版本，请刷新后重试")
    old = _latest_version(db, note)
    if old is None:
        raise HTTPException(status_code=404, detail="文件笔记不存在")
    filename = Path(file.filename or old.filename).name
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext != old.extension or ext not in ({"docx", "xlsx", "pdf", "md"} | IMAGE_EXTENSIONS):
        raise HTTPException(status_code=415, detail="保存时文件格式必须保持不变")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    _check_file(ext, content)
    persist_file(db, note, filename, content, ext)
    if ext in IMAGE_EXTENSIONS:
        from app.knowledge.markdown_images import invalidate_markdown_references

        invalidate_markdown_references(db, user_id, note.id)
    note.version += 1
    note.updated_at = datetime.now(timezone.utc)
    record_revision(db, note, db.scalars(select(NoteTag.tag_id).where(NoteTag.note_id == note.id, NoteTag.user_id == user_id)).all())
    record_event(db, user_id, AuditAction.UPDATE, AuditEntityType.NOTE, note.id, entity_version=note.version, details={"fields": ["content_kind"]})
    db.commit()
    db.refresh(note)
    return note_payload(db, note)


@router.get("/v1/notebooks/{notebook_id}/export")
def export_notebook(notebook_id: uuid.UUID, db: Db, user_id: UserId) -> Response:
    notebook = db.scalar(select(Notebook).where(Notebook.id == notebook_id, Notebook.user_id == user_id, Notebook.deleted_at.is_(None)))
    if notebook is None:
        raise HTTPException(status_code=404, detail="笔记本不存在")
    notes = db.scalars(select(Note).where(Note.user_id == user_id, Note.notebook_id == notebook_id, Note.deleted_at.is_(None)).order_by(Note.title, Note.id)).all()
    return _export_notes(db, notes, f"{safe_filename(notebook.name, 'notebook')}.zip")


def _export_notes(db: Db, notes: list[Note], filename: str) -> Response:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        used: set[str] = set()
        for note in notes:
            row = _latest_version(db, note)
            folder_name = db.scalar(select(Notebook.name).where(Notebook.id == note.notebook_id, Notebook.user_id == note.user_id)) if note.notebook_id else None
            folder = safe_filename(folder_name, "未分类") if folder_name else "未分类"
            if row:
                name = row.filename
                try:
                    content = _file_bytes(row)
                except HTTPException as exc:
                    raise HTTPException(status_code=502, detail="导出失败：有文件内容不可读取，请检查后重试") from exc
            else:
                name = f"{safe_filename(note.title)}.md"
                content = note.body_md.encode("utf-8")
            path_name = f"{folder}/{name}"
            if path_name in used:
                stem, suffix = os.path.splitext(name)
                ordinal = 2
                while f"{folder}/{stem} ({ordinal}){suffix}" in used:
                    ordinal += 1
                name = f"{stem} ({ordinal}){suffix}"
                path_name = f"{folder}/{name}"
            used.add(path_name)
            bundle.writestr(path_name, content)
    return download_response(archive.getvalue(), filename, "application/zip")


@router.get("/v1/files/{storage_key}")
def get_file_for_editor(storage_key: str, token: str, db: Db) -> Response:
    secret = os.getenv("PKM_FILE_LINK_SECRET", "")
    if not secret:
        raise HTTPException(status_code=503, detail="文件编辑链接未配置")
    try:
        claims = jwt.decode(token, secret, algorithms=["HS256"])
        if claims.get("key") != storage_key or claims.get("exp", 0) < int(datetime.now(timezone.utc).timestamp()):
            raise ValueError("expired")
    except (jwt.InvalidTokenError, ValueError):
        raise HTTPException(status_code=401, detail="文件访问链接已失效") from None
    row = db.scalar(select(NoteFileVersion).where(NoteFileVersion.storage_key == storage_key))
    if row is not None:
        return Response(_file_bytes(row), media_type=claims.get("media_type", row.media_type))
    if re.fullmatch(r"[0-9a-f-]{36}", storage_key):
        try:
            return Response(object_storage.get(storage_key), media_type=claims.get("media_type", "application/octet-stream"))
        except Exception:
            path = _file_path(storage_key)
            if path.is_file():
                return Response(path.read_bytes(), media_type=claims.get("media_type", "application/octet-stream"))
    raise HTTPException(status_code=404, detail="文件不存在")


@router.get("/v1/notes/{note_id}/editor-config")
def office_editor_config(note_id: uuid.UUID, db: Db, user_id: UserId) -> dict:
    note = owned_note(db, note_id, user_id)
    row = _latest_version(db, note)
    office_url = os.getenv("ONLYOFFICE_PUBLIC_URL", "").rstrip("/")
    app_url = os.getenv("PUBLIC_APP_URL", "http://api:8000").rstrip("/")
    if row is None or row.extension not in {"docx", "xlsx", "pdf"}:
        raise HTTPException(status_code=422, detail="此文件类型不支持文档编辑器")
    if not office_url or not os.getenv("PKM_FILE_LINK_SECRET"):
        return {"available": False, "message": "文档编辑服务尚未配置，可下载文件后编辑并重新上传"}
    expiry = int((datetime.now(timezone.utc) + timedelta(minutes=10)).timestamp())
    token = jwt.encode({"key": row.storage_key, "media_type": row.media_type, "exp": expiry}, os.environ["PKM_FILE_LINK_SECRET"], algorithm="HS256")
    file_url = f"{app_url}/v1/files/{row.storage_key}?token={token}"
    document_type = "cell" if row.extension == "xlsx" else "pdf" if row.extension == "pdf" else "word"
    callback_key = f"{note.id}-{note.version}-{row.version}-{row.sha256[:12]}"
    editor_config = {
        "documentType": document_type,
        "document": {
            # 加入笔记并发版本，避免笔记恢复或更新后 ONLYOFFICE 继续使用过期的编辑器缓存。
            "key": callback_key,
            "url": file_url,
            "fileType": row.extension,
            "title": row.filename,
            "permissions": {"edit": True, "download": True},
        },
        "editorConfig": {
            "mode": "edit",
            "lang": "zh",
            "region": "zh-CN",
            "customization": {"zoom": 80, "forcesave": True},
            "callbackUrl": f"{app_url}/v1/integrations/onlyoffice/callback/{note.id}?version={note.version}&file_version={row.version}&file_sha256={row.sha256}",
            "user": {"id": str(user_id), "name": "当前用户"},
        },
        "width": "100%",
        "height": "100%",
    }
    editor_token = jwt.encode(editor_config, os.getenv("ONLYOFFICE_JWT_SECRET", os.environ["PKM_FILE_LINK_SECRET"]), algorithm="HS256")
    return {"available": True, "api_url": f"{office_url}/web-apps/apps/api/documents/api.js", "token": editor_token, **editor_config}


@router.post("/v1/integrations/onlyoffice/callback/{note_id}")
async def onlyoffice_callback(
    note_id: uuid.UUID,
    request: Request,
    db: Db,
    request_body: dict,
    version: int = Query(ge=1),
    file_version: int = Query(ge=1),
    file_sha256: str = Query(pattern=r"^[0-9a-f]{64}$"),
) -> dict:
    secret = os.getenv("ONLYOFFICE_JWT_SECRET", os.getenv("PKM_FILE_LINK_SECRET", ""))
    note = db.scalar(select(Note).where(Note.id == note_id, Note.deleted_at.is_(None)))
    current_file = _latest_version(db, note) if note else None
    if (
        note is None
        or current_file is None
        or note.version != version
        or current_file.version != file_version
        or current_file.sha256 != file_sha256
    ):
        return {"error": 1}
    expected_key = f"{note.id}-{version}-{file_version}-{file_sha256[:12]}"
    if secret:
        token = request_body.pop("token", None)
        if not token:
            authorization = request.headers.get("authorization", "")
            token = authorization.removeprefix("Bearer ") if authorization.startswith("Bearer ") else authorization
        try:
            claims = jwt.decode(token or "", secret, algorithms=["HS256"])
            token_key = claims.get("payload", {}).get("key") or claims.get("key") or claims.get("document", {}).get("key")
            callback_key = request_body.get("key")
            if (
                not isinstance(callback_key, str)
                or callback_key != expected_key
                or token_key != callback_key
                or ("status" in claims and claims["status"] != request_body.get("status"))
            ):
                raise ValueError("document key mismatch")
        except jwt.InvalidTokenError:
            raise HTTPException(status_code=401, detail="文档回调签名无效") from None
        except ValueError:
            raise HTTPException(status_code=401, detail="文档回调版本无效") from None
    status = request_body.get("status")
    if status in {3, 7}:
        return {"error": 1}
    if status not in {2, 6}:
        return {"error": 0}
    source = _onlyoffice_download_url(request_body.get("url"))
    try:
        async with httpx.AsyncClient(timeout=15, follow_redirects=False) as client:
            async with client.stream("GET", source) as response:
                response.raise_for_status()
                content = bytearray()
                async for chunk in response.aiter_bytes():
                    content.extend(chunk)
                    if len(content) > MAX_UPLOAD_BYTES:
                        raise HTTPException(status_code=413, detail="文档服务返回文件超出大小限制")
    except httpx.HTTPError:
        raise HTTPException(status_code=502, detail="无法读取文档服务保存结果") from None
    note = db.scalar(select(Note).where(Note.id == note_id, Note.deleted_at.is_(None)).with_for_update())
    if note is None or note.version != version:
        return {"error": 1}
    row = _latest_version(db, note)
    if row is None or row.version != file_version or row.sha256 != file_sha256:
        return {"error": 1}
    _check_file(row.extension, content)
    persist_file(db, note, row.filename, content, row.extension)
    note.updated_at = datetime.now(timezone.utc)
    if status == 2:
        note.version += 1
        record_revision(db, note, db.scalars(select(NoteTag.tag_id).where(NoteTag.note_id == note.id, NoteTag.user_id == note.user_id)).all())
    # 文档服务回写：无登录态，归属笔记所有者，标记为系统行为。
    record_event(
        db, note.user_id, AuditAction.UPDATE, AuditEntityType.NOTE, note.id,
        entity_version=note.version,
        details={"changed": ["file_version"], "source": "onlyoffice", "document_status": status},
        actor_type=AuditActor.SYSTEM,
    )
    db.commit()
    return {"error": 0}


class LinkDraftCreate(BaseModel):
    url: HttpUrl
    notebook_id: uuid.UUID | None = None
    pasted_text: str = Field(default="", max_length=100_000)


class LinkDraftUpdate(BaseModel):
    title: str | None = Field(default=None, min_length=1, max_length=240)
    body_md: str | None = Field(default=None, max_length=100_000)
    pasted_text: str | None = Field(default=None, max_length=100_000)
    notebook_id: uuid.UUID | None = None


class LinkDraftPublish(BaseModel):
    notebook_id: uuid.UUID | None = None
    tag_ids: list[uuid.UUID] = Field(default_factory=list, max_length=100)


LinkDraftAnalysis = ContentAnalysis


def _validate_url_syntax(url: str) -> None:
    if len(url) > 2048:
        raise HTTPException(status_code=422, detail="链接超过 2048 个字符")
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise HTTPException(status_code=422, detail="只支持公开的 HTTP/HTTPS 网页链接")
    try:
        port = parsed.port
    except ValueError:
        raise HTTPException(status_code=422, detail="网页链接端口无效") from None
    allowed_port = 80 if parsed.scheme == "http" else 443
    if port is not None and port != allowed_port:
        raise HTTPException(status_code=422, detail="出于安全原因，只能抓取 HTTP/HTTPS 标准端口")


def canonicalize_source_url(url: str) -> str:
    """Normalize only URL components declared equivalent by the product policy."""
    _validate_url_syntax(url)
    parsed = urlsplit(url)
    scheme = parsed.scheme.lower()
    hostname = parsed.hostname or ""
    try:
        hostname = hostname.encode("idna").decode("ascii").lower()
        port = parsed.port
    except (UnicodeError, ValueError):
        raise HTTPException(status_code=422, detail="网页链接无效") from None
    if ":" in hostname and not hostname.startswith("["):
        hostname = f"[{hostname}]"
    default_port = 443 if scheme == "https" else 80
    authority = hostname if port in {None, default_port} else f"{hostname}:{port}"
    return urlunsplit((scheme, authority, parsed.path, parsed.query, ""))


def _resolve_public_addresses(url: str) -> list[str]:
    _validate_url_syntax(url)
    parsed = urlparse(url)
    try:
        results = socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)
    except (OSError, ValueError):
        raise HTTPException(status_code=422, detail="网页域名无法解析") from None
    addresses = list(dict.fromkeys(item[4][0] for item in results))
    try:
        all_public = bool(addresses) and all(ipaddress.ip_address(address.split("%", 1)[0]).is_global for address in addresses)
    except ValueError:
        all_public = False
    if not all_public:
        raise HTTPException(status_code=422, detail="出于安全原因，不能抓取内网或保留地址")
    return addresses


class _PinnedHTTPConnection(http.client.HTTPConnection):
    def __init__(self, host: str, port: int, addresses: list[str], timeout: float):
        super().__init__(host, port, timeout=timeout)
        self.addresses = addresses

    def connect(self) -> None:
        errors: list[OSError] = []
        for address in self.addresses:
            try:
                self.sock = socket.create_connection((address, self.port), self.timeout, self.source_address)
                return
            except OSError as exc:
                errors.append(exc)
        if errors:
            raise errors[-1]
        raise OSError("no validated address available")


class _PinnedHTTPSConnection(_PinnedHTTPConnection):
    def connect(self) -> None:
        errors: list[OSError] = []
        context = ssl.create_default_context()
        for address in self.addresses:
            sock = None
            try:
                sock = socket.create_connection((address, self.port), self.timeout, self.source_address)
                self.sock = context.wrap_socket(sock, server_hostname=self.host)
                return
            except OSError as exc:
                errors.append(exc)
                if sock is not None:
                    sock.close()
        if errors:
            raise errors[-1]
        raise OSError("no validated address available")


def _fetch_pinned_response(url: str) -> tuple[int, str, str | None, bytes]:
    parsed = urlsplit(url)
    addresses = _resolve_public_addresses(url)
    port = parsed.port or (443 if parsed.scheme == "https" else 80)
    host = (parsed.hostname or "").encode("idna").decode("ascii")
    connection_type = _PinnedHTTPSConnection if parsed.scheme == "https" else _PinnedHTTPConnection
    connection = connection_type(host, port, addresses, timeout=4)
    target = parsed.path or "/"
    if parsed.query:
        target += f"?{parsed.query}"
    host_header = host
    if ":" in host and not host.startswith("["):
        host_header = f"[{host}]"
    if parsed.port is not None:
        host_header = f"{host_header}:{parsed.port}"
    try:
        connection.request("GET", target, headers={"Host": host_header, "User-Agent": "MapleNotes/1.0", "Accept-Encoding": "identity", "Connection": "close"})
        response = connection.getresponse()
        content_type = response.getheader("content-type", "")
        location = response.getheader("location")
        status = response.status
        if status in {301, 302, 303, 307, 308} or status == 404:
            return status, content_type, location, b""
        if status < 200 or status >= 300:
            raise HTTPException(status_code=502, detail="网页抓取失败，可手动粘贴正文")
        raw = bytearray()
        while chunk := response.read(64 * 1024):
            raw.extend(chunk)
            if len(raw) > MAX_WEB_BYTES:
                raise HTTPException(status_code=413, detail="网页正文超过 3 MB 限制")
        return status, content_type, location, bytes(raw)
    finally:
        connection.close()


def fetch_public_page_source(url: str, *, _check_robots: bool = True) -> tuple[str, str, bytes]:
    """Fetch a bounded public-page response without parsing its HTML content."""
    _resolve_public_addresses(url)
    if _check_robots:
        parsed = urlparse(url)
        origin = f"{parsed.scheme}://{parsed.netloc}"
        robots_url = f"{origin}/robots.txt"
        _, robots_type, robots_raw = fetch_public_page_source(robots_url, _check_robots=False)
        if robots_raw:
            if "text/plain" not in robots_type:
                raise HTTPException(status_code=415, detail="网站抓取规则不是可识别的纯文本")
            policy = RobotFileParser(robots_url)
            policy.parse(robots_raw.decode("utf-8", errors="replace").splitlines())
            if not policy.can_fetch("MapleNotes/1.0", url):
                raise HTTPException(status_code=403, detail="网站 robots.txt 不允许抓取该页面")
    current = url
    for _ in range(4):
        try:
            status, content_type, location, raw = _fetch_pinned_response(current)
            if status in {301, 302, 303, 307, 308}:
                if not location:
                    raise HTTPException(status_code=422, detail="网页重定向地址无效")
                current = urljoin(current, location)
                continue
            if status == 404:
                return current, content_type or "text/plain", b""
            if "html" not in content_type and "text/plain" not in content_type:
                raise HTTPException(status_code=415, detail="链接目标不是可识别的网页")
            return current, content_type, raw
        except (socket.timeout, TimeoutError):
            raise HTTPException(status_code=504, detail="网页请求超时") from None
        except (OSError, http.client.HTTPException, ssl.SSLError):
            raise HTTPException(status_code=502, detail="网页抓取失败，可手动粘贴正文") from None
    raise HTTPException(status_code=422, detail="网页重定向次数过多")


_NON_TEXT_TAGS = {
    "script", "style", "nav", "footer", "header", "noscript", "svg", "img",
    "picture", "video", "audio", "source", "iframe", "object", "embed", "canvas",
}


def _markdown_inline(node) -> str:
    if isinstance(node, NavigableString):
        return str(node)
    if not isinstance(node, Tag) or node.name in _NON_TEXT_TAGS:
        return ""

    name = node.name.lower()
    if name == "br":
        return "\n"
    if name == "code":
        value = node.get_text("", strip=False).strip()
        fence = "`" * max(1, max((len(run) for run in re.findall(r"`+", value)), default=0) + 1)
        return f"{fence}{value}{fence}"
    value = "".join(_markdown_inline(child) for child in node.children)
    if name in {"strong", "b"}:
        return f"**{value.strip()}**" if value.strip() else ""
    if name in {"em", "i"}:
        return f"*{value.strip()}*" if value.strip() else ""
    if name in {"del", "s", "strike"}:
        return f"~~{value.strip()}~~" if value.strip() else ""
    if name == "a":
        href = (node.get("href") or "").strip()
        scheme = urlparse(href).scheme.lower()
        safe_link = not scheme or scheme in {"http", "https", "mailto"}
        return f"[{value.strip()}]({href})" if href and safe_link and value.strip() else value
    if name in {"p", "div", "section", "article", "main", "span"}:
        return value
    return value


def _markdown_list(node: Tag, depth: int = 0) -> str:
    ordered = node.name.lower() == "ol"
    start = int(node.get("start", 1)) if ordered and str(node.get("start", "1")).isdigit() else 1
    indent = "  " * depth
    rows = []
    item_index = 0
    for item in node.find_all("li", recursive=False):
        item_index += 1
        content_parts = []
        nested_lists = []
        for child in item.children:
            if isinstance(child, Tag) and child.name.lower() in {"ol", "ul"}:
                nested_lists.append(_markdown_list(child, depth + 1))
            else:
                content_parts.append(_markdown_inline(child).strip())
        content = " ".join(part for part in content_parts if part).strip()
        number = item.get("value") if ordered else None
        marker = f"{number or start + item_index - 1}. " if ordered else "- "
        lines = content.splitlines() or [""]
        rows.append(indent + marker + lines[0])
        rows.extend(indent + "  " + line for line in lines[1:])
        rows.extend(nested_lists)
    return "\n".join(rows)


def _markdown_table(node: Tag) -> str:
    rows = []
    for row in node.find_all("tr"):
        cells = row.find_all(["th", "td"], recursive=False)
        if cells:
            rows.append([_markdown_inline(cell).strip().replace("|", "\\|").replace("\n", " <br> ") for cell in cells])
    if not rows:
        return ""
    width = max(map(len, rows))
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = "| " + " | ".join(rows[0]) + " |"
    divider = "| " + " | ".join("---" for _ in range(width)) + " |"
    body = ["| " + " | ".join(row) + " |" for row in rows[1:]]
    return "\n".join([header, divider, *body])


def _html_to_markdown(node) -> str:
    if isinstance(node, NavigableString):
        return str(node)
    if not isinstance(node, Tag) or node.name in _NON_TEXT_TAGS:
        return ""
    name = node.name.lower()
    if name in {"ul", "ol"}:
        return _markdown_list(node)
    if name == "table":
        return _markdown_table(node)
    if name == "pre":
        code = node.get_text("", strip=False).strip("\n")
        classes = " ".join(node.get("class", []))
        child = node.find("code")
        if child:
            classes += " " + " ".join(child.get("class", []))
        match = re.search(r"(?:language|lang)-([\w+#.-]+)", classes)
        language = match.group(1) if match else ""
        fence_size = max(3, max((len(run) for run in re.findall(r"`+", code)), default=0) + 1)
        fence = "`" * fence_size
        return f"\n\n{fence}{language}\n{code}\n{fence}\n\n"
    value = "".join(_html_to_markdown(child) for child in node.children)
    if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        return f"\n\n{'#' * int(name[1])} {value.strip()}\n\n"
    if name == "blockquote":
        quoted = "\n".join("> " + line if line else ">" for line in value.strip().splitlines())
        return f"\n\n{quoted}\n\n"
    if name == "li":
        return value
    if name == "hr":
        return "\n\n---\n\n"
    if name in {"p", "div", "section", "article", "main", "body"}:
        return f"\n\n{value.strip()}\n\n" if value.strip() else ""
    return _markdown_inline(node)


def parse_public_page_source(content_type: str, raw: bytes) -> tuple[str, str]:
    """Extract bounded text while preserving common Markdown structures; never extracts media."""
    if "html" not in content_type and "text/plain" not in content_type:
        raise HTTPException(status_code=415, detail="链接目标不是可识别的网页")
    soup = BeautifulSoup(raw, "html.parser")
    title = (soup.title.get_text(" ", strip=True) if soup.title else "")[:240]
    if "text/plain" in content_type:
        text = raw.decode(soup.original_encoding or "utf-8", errors="replace").strip()
    else:
        for element in soup(_NON_TEXT_TAGS):
            element.decompose()
        main = soup.find("article") or soup.find("main") or soup.body or soup
        text = _html_to_markdown(main)
        text = re.sub(r"[ \t]+\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return title, text[:100_000]


def fetch_public_page(url: str) -> tuple[str, str]:
    """Compatibility wrapper used by the background link-draft worker."""
    final_url, content_type, raw = fetch_public_page_source(url)
    return parse_public_page_source(content_type, raw)


def draft_json(draft: LinkDraft) -> dict:
    return {"id": str(draft.id), "source_url": draft.source_url, "title": draft.title, "snapshot_text": draft.snapshot_text, "body_md": draft.body_md, "fetch_status": draft.fetch_status, "fetch_error": draft.fetch_error, "status": draft.status, "notebook_id": str(draft.notebook_id) if draft.notebook_id else None, "created_at": draft.created_at.isoformat(), "updated_at": draft.updated_at.isoformat()}


@router.post("/v1/link-drafts", status_code=201)
def create_link_draft(body: LinkDraftCreate, db: Db, user_id: UserId) -> dict:
    url = canonicalize_source_url(str(body.url))
    _resolve_public_addresses(url)
    existing_note = next((
        row for row in db.scalars(select(Note).where(Note.user_id == user_id, Note.source_url.is_not(None), Note.deleted_at.is_(None))).all()
        if canonicalize_source_url(row.source_url) == url
    ), None)
    if existing_note:
        raise HTTPException(status_code=409, detail=f"该链接已导入为笔记：{existing_note.id}")
    existing = next((
        row for row in db.scalars(select(LinkDraft).where(LinkDraft.user_id == user_id, LinkDraft.status != "published")).all()
        if canonicalize_source_url(row.source_url) == url
    ), None)
    if existing:
        raise HTTPException(status_code=409, detail=f"该链接已有草稿：{existing.id}")
    validate_categories(db, user_id, body.notebook_id, [])
    title = urlparse(url).hostname or "网页草稿"
    text = body.pasted_text
    draft = LinkDraft(
        user_id=user_id, source_url=url, title=title, snapshot_text=text, body_md=text,
        fetch_status="ready" if text else "pending", fetch_error=None,
        notebook_id=body.notebook_id,
    )
    db.add(draft)
    db.flush()
    record_event(db, user_id, AuditAction.CREATE, AuditEntityType.LINK_DRAFT, draft.id,
                 details={"changed": ["source_url", "title", "notebook_id"]})
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        duplicate = next((
            row for row in db.scalars(select(LinkDraft).where(LinkDraft.user_id == user_id, LinkDraft.status != "published")).all()
            if canonicalize_source_url(row.source_url) == url
        ), None)
        if duplicate:
            raise HTTPException(status_code=409, detail=f"该链接已有草稿：{duplicate.id}") from None
        raise
    db.refresh(draft)
    return draft_json(draft)


@router.get("/v1/link-drafts")
def list_link_drafts(db: Db, user_id: UserId) -> dict:
    rows = db.scalars(select(LinkDraft).where(LinkDraft.user_id == user_id, LinkDraft.status == "draft").order_by(LinkDraft.updated_at.desc())).all()
    return {"items": [draft_json(row) for row in rows]}


@router.get("/v1/link-drafts/{draft_id}")
def get_link_draft(draft_id: uuid.UUID, db: Db, user_id: UserId) -> dict:
    draft = db.scalar(select(LinkDraft).where(LinkDraft.id == draft_id, LinkDraft.user_id == user_id, LinkDraft.status == "draft"))
    if draft is None:
        raise HTTPException(status_code=404, detail="网页草稿不存在")
    return draft_json(draft)


@router.patch("/v1/link-drafts/{draft_id}")
def update_link_draft(draft_id: uuid.UUID, body: LinkDraftUpdate, db: Db, user_id: UserId) -> dict:
    draft = db.scalar(select(LinkDraft).where(LinkDraft.id == draft_id, LinkDraft.user_id == user_id, LinkDraft.status == "draft").with_for_update())
    if draft is None:
        raise HTTPException(status_code=404, detail="网页草稿不存在")
    if "notebook_id" in body.model_fields_set:
        validate_categories(db, user_id, body.notebook_id, [])
    for key in ("title", "body_md", "notebook_id"):
        if key in body.model_fields_set:
            setattr(draft, key, getattr(body, key))
    if "body_md" in body.model_fields_set and draft.fetch_status in {"pending", "processing"}:
        draft.fetch_status = "ready"
        draft.fetch_error = None
        if not draft.snapshot_text.strip():
            draft.snapshot_text = body.body_md or ""
    if body.pasted_text is not None:
        draft.snapshot_text = body.pasted_text
        draft.body_md = body.pasted_text
        draft.fetch_status = "ready"
        draft.fetch_error = None
    draft.updated_at = datetime.now(timezone.utc)
    record_event(db, user_id, AuditAction.UPDATE, AuditEntityType.LINK_DRAFT, draft.id,
                 details={"changed": sorted(body.model_fields_set)})
    db.commit()
    db.refresh(draft)
    return draft_json(draft)


@router.post("/v1/link-drafts/{draft_id}/rewrite")
def rewrite_link_draft(draft_id: uuid.UUID, db: Db, user_id: UserId) -> dict:
    from app.assistant.model_connection import chat_model, decrypt_key, require_connection
    draft = db.scalar(select(LinkDraft).where(LinkDraft.id == draft_id, LinkDraft.user_id == user_id, LinkDraft.status == "draft"))
    if draft is None:
        raise HTTPException(status_code=404, detail="网页草稿不存在")
    if not draft.snapshot_text.strip():
        raise HTTPException(status_code=422, detail="没有可改写的网页正文，请先粘贴正文")
    check_limit(user_id, "assistant", limit=10)
    row = require_connection(db, user_id)
    model = chat_model(decrypt_key(row), row.model_name)
    try:
        suggestion = generate_link_note_content(
            model,
            source_title=draft.title,
            source_url=draft.source_url,
            content=draft.snapshot_text,
        )
    except Exception:
        raise HTTPException(status_code=502, detail="Agent 改写暂不可用") from None
    return {**draft_json(draft), "rewrite_suggestion": suggestion}


@router.post("/v1/link-drafts/{draft_id}/analyze")
def analyze_link_draft(draft_id: uuid.UUID, db: Db, user_id: UserId) -> dict:
    from app.assistant.model_connection import chat_model, decrypt_key, require_connection

    draft = db.scalar(select(LinkDraft).where(LinkDraft.id == draft_id, LinkDraft.user_id == user_id, LinkDraft.status == "draft"))
    if draft is None:
        raise HTTPException(status_code=404, detail="网页草稿不存在")
    if not draft.snapshot_text.strip():
        raise HTTPException(status_code=422, detail="没有可分析的网页正文，请先粘贴正文")
    check_limit(user_id, "assistant", limit=10)
    row = require_connection(db, user_id)
    model = chat_model(decrypt_key(row), row.model_name, max_tokens=1400)
    try:
        parsed = analyze_content(model, draft.snapshot_text[:30_000], source_label="网页正文")
    except Exception:
        raise HTTPException(status_code=502, detail="建议分析暂不可用，请稍后重试") from None
    return parsed.model_dump()


@router.post("/v1/link-drafts/{draft_id}/publish", status_code=201)
def publish_link_draft(draft_id: uuid.UUID, body: LinkDraftPublish, db: Db, user_id: UserId) -> dict:
    draft = db.scalar(select(LinkDraft).where(LinkDraft.id == draft_id, LinkDraft.user_id == user_id, LinkDraft.status == "draft").with_for_update())
    if draft is None:
        raise HTTPException(status_code=404, detail="网页草稿不存在")
    if not draft.body_md.strip():
        raise HTTPException(status_code=422, detail="草稿正文为空，无法发布")
    validate_categories(db, user_id, body.notebook_id, body.tag_ids)
    canonical_url = canonicalize_source_url(draft.source_url)
    existing_note = next((
        row for row in db.scalars(select(Note).where(Note.user_id == user_id, Note.source_url.is_not(None), Note.deleted_at.is_(None))).all()
        if canonicalize_source_url(row.source_url) == canonical_url
    ), None)
    if existing_note:
        raise HTTPException(status_code=409, detail=f"该链接已导入为笔记：{existing_note.id}")
    existing_draft = next((
        row for row in db.scalars(select(LinkDraft).where(LinkDraft.user_id == user_id, LinkDraft.status == "draft", LinkDraft.id != draft.id)).all()
        if canonicalize_source_url(row.source_url) == canonical_url
    ), None)
    if existing_draft:
        raise HTTPException(status_code=409, detail=f"该链接已有草稿：{existing_draft.id}")
    note = Note(user_id=user_id, notebook_id=body.notebook_id, title=draft.title, body_md=draft.body_md, version=1, content_version=1, content_kind="markdown", source_url=canonical_url)
    db.add(note)
    db.flush()
    for tag_id in dict.fromkeys(body.tag_ids):
        db.add(NoteTag(note_id=note.id, tag_id=tag_id, user_id=user_id))
    queue_index(db, note)
    record_revision(db, note, list(dict.fromkeys(body.tag_ids)))
    record_event(db, user_id, AuditAction.CREATE, AuditEntityType.NOTE, note.id, entity_version=note.version, details={"fields": ["source_url"]})
    draft.status = "published"
    draft.updated_at = datetime.now(timezone.utc)
    db.commit()
    db.refresh(note)
    from app.knowledge.notes import note_json
    return note_json(db, note)


@router.delete("/v1/link-drafts/{draft_id}", status_code=204)
def delete_link_draft(draft_id: uuid.UUID, db: Db, user_id: UserId) -> None:
    draft = db.scalar(select(LinkDraft).where(LinkDraft.id == draft_id, LinkDraft.user_id == user_id, LinkDraft.status == "draft").with_for_update())
    if draft is None:
        raise HTTPException(status_code=404, detail="网页草稿不存在")
    record_event(db, user_id, AuditAction.DELETE, AuditEntityType.LINK_DRAFT, draft.id)
    db.delete(draft)
    db.commit()
